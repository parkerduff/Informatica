"""Generate a self-contained XLSX workbook covering the converted jobs:

  * Overview            - per-job engine/source/target/mapping/metrics/verdict
  * Source Fields       - every SOURCEFIELD (offset/length/precision/scale/null)
  * Target Fields       - every TARGETFIELD
  * Transformations     - ordered transformation chain per mapping
  * Transform Expr      - every TRANSFORMFIELD expression
  * Reconciliation      - functional + performance recon results per job
  * <FOLDER>            - the actual converted output rows (all fields) per job

Run:  python3 migration/reports/generate_xlsx.py [--mode functional|performance]
Reads specs from migration/spec, results from migration/reports/_results_*.json,
and converted output CSVs from migration/local/out/<mode>/<folder>/.
"""
from __future__ import annotations

import argparse
import csv
import json
import os
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import sys as _sys
_sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..")))
from migration.jobs import configs  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
SPEC_DIR = os.path.join(ROOT, "migration", "spec")
OUT_DIR = os.path.join(ROOT, "migration", "local", "out")
REPORT_DIR = os.path.join(ROOT, "migration", "reports")

FOLDERS = ["Pseudossn", "EHRP2BIIS_UPDATE", "CPM_NIH", "CPM_OIG", "CPM_CDC",
           "FDA_Leave", "Pay_Calendar", "COMPTIME"]
ENGINE = {f: "Glue PySpark" for f in FOLDERS}
ENGINE["Pay_Calendar"] = "Glue Python Shell"
ENGINE["COMPTIME"] = "Glue Python Shell"

# --- styling -------------------------------------------------------------
HDR_FILL = PatternFill("solid", fgColor="1E293B")
HDR_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(color="0F172A", bold=True, size=14)
ACCENT_FONT = Font(color="0E7490", bold=True)
PASS_FILL = PatternFill("solid", fgColor="D1FAE5")
FAIL_FILL = PatternFill("solid", fgColor="FEE2E2")
ZEBRA = PatternFill("solid", fgColor="F1F5F9")
THIN = Side(style="thin", color="CBD5E1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="top", wrap_text=True)
TOP = Alignment(vertical="top")


def _primary_target(folder: str) -> str:
    """The actual converted primary target (matches the output CSV filename)."""
    try:
        return configs.build(folder).target.name
    except Exception:
        sp = _load_spec(folder)
        return sp["targets"][0]["name"] if sp["targets"] else folder


def _load_spec(folder: str) -> Dict[str, Any]:
    with open(os.path.join(SPEC_DIR, f"{folder}.json")) as fh:
        return json.load(fh)


def _load_results(mode: str) -> Optional[Dict[str, Any]]:
    p = os.path.join(REPORT_DIR, f"_results_{mode}.json")
    if os.path.exists(p):
        with open(p) as fh:
            return json.load(fh)
    return None


def _header(ws, headers: List[str], row: int = 1) -> None:
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _autosize(ws, max_width: int = 60) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[letter].width = min(max(width + 2, 10), max_width)


def _write_rows(ws, rows: List[List[Any]], start: int = 2, zebra: bool = True) -> None:
    for i, r in enumerate(rows):
        rr = start + i
        for c, v in enumerate(r, 1):
            cell = ws.cell(row=rr, column=c, value=v)
            cell.border = BORDER
            cell.alignment = TOP
        if zebra and i % 2 == 1:
            for c in range(1, len(r) + 1):
                ws.cell(row=rr, column=c).fill = ZEBRA


# --- sheets --------------------------------------------------------------
def sheet_overview(wb: Workbook, specs: Dict[str, Any],
                   res: Dict[str, Optional[Dict[str, Any]]]) -> None:
    ws = wb.active
    ws.title = "Overview"
    ws.cell(row=1, column=1, value="Informatica PowerCenter \u2192 AWS Glue \u2014 Converted Jobs").font = TITLE_FONT
    headers = ["Folder", "Engine", "Mapping", "Primary source", "Primary target",
               "Source fields", "Target fields", "Transforms",
               "Func rows", "Func verdict", "Perf rows", "Perf verdict"]
    _header(ws, headers, row=3)
    rows = []
    fres, pres = res.get("functional"), res.get("performance")
    for f in FOLDERS:
        sp = specs[f]
        src = sp["sources"][0]["name"] if sp["sources"] else ""
        tgt = _primary_target(f)
        nsf = sum(len(s.get("fields", [])) for s in sp["sources"])
        ntf = sum(len(t.get("fields", [])) for t in sp["targets"])
        ntr = sum(len(m.get("transformations", [])) for m in sp["mappings"])
        fr = (fres or {}).get("results", {}).get(f, {})
        pr = (pres or {}).get("results", {}).get(f, {})
        rows.append([
            f, ENGINE[f], sp["mappings"][0]["name"] if sp["mappings"] else "", src, tgt,
            nsf, ntf, ntr,
            fr.get("conv_metrics", {}).get("detail_rows", ""), fr.get("verdict", ""),
            pr.get("conv_metrics", {}).get("detail_rows", ""), pr.get("verdict", ""),
        ])
    _write_rows(ws, rows, start=4)
    for i in range(len(rows)):
        for col in (10, 12):
            cell = ws.cell(row=4 + i, column=col)
            if cell.value == "PASS":
                cell.fill = PASS_FILL
                cell.font = Font(color="065F46", bold=True)
            elif cell.value == "FAIL":
                cell.fill = FAIL_FILL
                cell.font = Font(color="991B1B", bold=True)
    _autosize(ws)


def sheet_source_fields(wb: Workbook, specs: Dict[str, Any]) -> None:
    ws = wb.create_sheet("Source Fields")
    _header(ws, ["Folder", "Source", "Field #", "Name", "Datatype", "Precision",
                 "Scale", "Phys offset", "Phys length", "Nullable", "Key type"])
    rows = []
    for f in FOLDERS:
        for s in specs[f]["sources"]:
            for fld in s.get("fields", []):
                rows.append([f, s["name"], fld.get("fieldnumber"), fld.get("name"),
                             fld.get("datatype"), fld.get("precision"), fld.get("scale"),
                             fld.get("physicaloffset"), fld.get("physicallength"),
                             fld.get("nullable"), fld.get("keytype")])
    _write_rows(ws, rows)
    _autosize(ws)


def sheet_target_fields(wb: Workbook, specs: Dict[str, Any]) -> None:
    ws = wb.create_sheet("Target Fields")
    _header(ws, ["Folder", "Target", "Field #", "Name", "Datatype", "Precision",
                 "Scale", "Nullable", "Key type"])
    rows = []
    for f in FOLDERS:
        for t in specs[f]["targets"]:
            for i, fld in enumerate(t.get("fields", []), 1):
                rows.append([f, t["name"], fld.get("fieldnumber", i), fld.get("name"),
                             fld.get("datatype"), fld.get("precision"), fld.get("scale"),
                             fld.get("nullable"), fld.get("keytype")])
    _write_rows(ws, rows)
    _autosize(ws)


def sheet_transformations(wb: Workbook, specs: Dict[str, Any]) -> None:
    ws = wb.create_sheet("Transformations")
    _header(ws, ["Folder", "Mapping", "Order", "Transformation", "Type", "Ports"])
    rows = []
    for f in FOLDERS:
        for m in specs[f]["mappings"]:
            for i, tr in enumerate(m.get("transformations", []), 1):
                rows.append([f, m["name"], i, tr.get("name"), tr.get("type"),
                             len(tr.get("fields", []))])
    _write_rows(ws, rows)
    _autosize(ws)


def sheet_expressions(wb: Workbook, specs: Dict[str, Any]) -> None:
    ws = wb.create_sheet("Transform Expr")
    _header(ws, ["Folder", "Mapping", "Transformation", "Type", "Port", "Port type",
                 "Datatype", "Prec", "Scale", "Expression"])
    rows = []
    for f in FOLDERS:
        for m in specs[f]["mappings"]:
            for tr in m.get("transformations", []):
                for fld in tr.get("fields", []):
                    expr = fld.get("expression") or ""
                    if not expr:
                        continue
                    rows.append([f, m["name"], tr.get("name"), tr.get("type"),
                                 fld.get("name"), fld.get("porttype"),
                                 fld.get("datatype"), fld.get("precision"),
                                 fld.get("scale"), expr])
    _write_rows(ws, rows)
    _autosize(ws)
    ws.column_dimensions[get_column_letter(10)].width = 80
    for r in range(2, len(rows) + 2):
        ws.cell(row=r, column=10).alignment = WRAP


def sheet_reconciliation(wb: Workbook, res: Dict[str, Optional[Dict[str, Any]]]) -> None:
    ws = wb.create_sheet("Reconciliation")
    _header(ws, ["Mode", "Folder", "Engine", "Verdict", "Conv rows", "Base rows",
                 "Error rows", "Cells compared", "Cell mismatches", "Null mismatches",
                 "Field match %", "Error parity", "Counter parity",
                 "Runtime s", "Rows/sec"])
    rows = []
    for mode in ("functional", "performance"):
        data = res.get(mode)
        if not data:
            continue
        for f in FOLDERS:
            r = data.get("results", {}).get(f)
            if not r:
                continue
            cm = r.get("conv_metrics", {})
            fld = r.get("field", {})
            rc = r.get("row_counts", {})
            rows.append([
                mode, f, "Glue PySpark" if cm.get("engine") == "pyspark" else "Glue Python Shell",
                r.get("verdict"),
                rc.get("main", {}).get("conv"), rc.get("main", {}).get("base"),
                rc.get("error", {}).get("conv"),
                fld.get("total_cells"), fld.get("cell_mismatches"),
                fld.get("null_mismatches"),
                round(fld.get("field_match_rate", 0) * 100, 4),
                "yes" if r.get("error_parity", {}).get("keys_match") else "no",
                "yes" if r.get("counter_parity", {}).get("match") else "no",
                cm.get("runtime_sec"), cm.get("rows_per_sec"),
            ])
    _write_rows(ws, rows)
    for i in range(len(rows)):
        cell = ws.cell(row=2 + i, column=4)
        if cell.value == "PASS":
            cell.fill = PASS_FILL
            cell.font = Font(color="065F46", bold=True)
        elif cell.value == "FAIL":
            cell.fill = FAIL_FILL
            cell.font = Font(color="991B1B", bold=True)
    _autosize(ws)


def _read_csv(path: str, limit: Optional[int] = None):
    with open(path, newline="") as fh:
        r = csv.reader(fh)
        header = next(r, [])
        rows = []
        for i, row in enumerate(r):
            if limit is not None and i >= limit:
                break
            rows.append(row)
    return header, rows


def sheet_job_data(wb: Workbook, specs: Dict[str, Any], mode: str,
                   max_rows: Optional[int]) -> int:
    written = 0
    for f in FOLDERS:
        tgt = _primary_target(f)
        csv_path = os.path.join(OUT_DIR, mode, f, f"{tgt}.csv")
        if not os.path.exists(csv_path):
            continue
        header, rows = _read_csv(csv_path, limit=max_rows)
        title = f[:31]
        ws = wb.create_sheet(title)
        _header(ws, header)
        _write_rows(ws, rows)
        _autosize(ws, max_width=40)
        written += 1
    return written


def main(argv: Optional[List[str]] = None) -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--mode", choices=["functional", "performance"], default="functional",
                    help="which run's converted output rows to embed as data sheets")
    ap.add_argument("--max-rows", type=int, default=1000,
                    help="cap converted-output rows per job sheet (0 = all)")
    ap.add_argument("--out", default=os.path.join(REPORT_DIR, "converted_jobs_report.xlsx"))
    args = ap.parse_args(argv)
    max_rows = None if args.max_rows == 0 else args.max_rows

    specs = {f: _load_spec(f) for f in FOLDERS}
    res = {"functional": _load_results("functional"),
           "performance": _load_results("performance")}

    wb = Workbook()
    sheet_overview(wb, specs, res)
    sheet_source_fields(wb, specs)
    sheet_target_fields(wb, specs)
    sheet_transformations(wb, specs)
    sheet_expressions(wb, specs)
    sheet_reconciliation(wb, res)
    ndata = sheet_job_data(wb, specs, args.mode, max_rows)

    wb.save(args.out)
    print(f"[xlsx] wrote {args.out}  ({len(wb.sheetnames)} sheets; "
          f"{ndata} job-data sheets from '{args.mode}' run)")


if __name__ == "__main__":
    main()
